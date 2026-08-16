import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def create_testing_report():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)
        
    # Styles
    styles = doc.styles
    style_normal = styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(13)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)
    
    style_h2 = styles['Heading 2']
    style_h2.font.name = 'Times New Roman'
    style_h2.font.size = Pt(14)
    style_h2.font.bold = True
    style_h2.font.color.rgb = RGBColor(0, 0, 0)
    
    style_h3 = styles['Heading 3']
    style_h3.font.name = 'Times New Roman'
    style_h3.font.size = Pt(13)
    style_h3.font.bold = True
    style_h3.font.italic = True
    style_h3.font.color.rgb = RGBColor(0, 0, 0)

    # Content
    doc.add_heading('PHẦN 4: KIỂM THỬ (TESTING)', level=1).style.font.name = 'Times New Roman'
    
    # 4.1
    doc.add_heading('4.1. Chiến lược kiểm thử (Testing Strategy)', level=2)
    doc.add_paragraph('Các loại kiểm thử được áp dụng trong dự án ExpenseAI bao gồm:')
    ul = ['Unit Test – Kiểm thử đơn vị (cho backend): Kiểm tra các hàm logic độc lập như tính tổng thu/chi, mã hóa mật khẩu.',
          'Integration Test – Kiểm thử tích hợp (API ↔ Database): Xác minh sự kết nối chính xác giữa các endpoint FastAPI và PostgreSQL.',
          'UI Test – Kiểm thử giao diện (thủ công): Đảm bảo giao diện hiển thị đúng trên các kích thước màn hình khác nhau.',
          'End-to-End Test – Kiểm thử luồng chính: Kiểm tra kịch bản xuyên suốt (Đăng nhập → Thêm giao dịch → Xem báo cáo).',
          'Security Test – Kiểm thử bảo mật: Kiểm tra các rủi ro bảo mật (XSS, SQL Injection, JWT timeout).',
          'Performance Test – Kiểm thử hiệu năng: Kiểm tra tốc độ tải trang và phản hồi API khi có khối lượng dữ liệu lớn.']
    for item in ul:
        doc.add_paragraph(item, style='List Bullet')

    # 4.2
    doc.add_heading('4.2. Danh sách Test Case chi tiết (58 test cases)', level=2)
    
    test_cases = [
        # Auth (8)
        ("Auth", "Đăng nhập thành công", "Tài khoản hợp lệ", "Nhập email/pass đúng", "Chuyển hướng Dashboard", "P0", "PASS"),
        ("Auth", "Đăng nhập sai mật khẩu", "Tài khoản hợp lệ", "Nhập pass sai", "Báo lỗi sai mật khẩu", "P0", "PASS"),
        ("Auth", "Đăng nhập sai email", "Email không tồn tại", "Nhập email sai", "Báo lỗi tài khoản không tồn tại", "P0", "PASS"),
        ("Auth", "Đăng ký thành công", "Email chưa tồn tại", "Nhập thông tin hợp lệ", "Tạo tài khoản, chuyển hướng Login", "P0", "PASS"),
        ("Auth", "Đăng ký email đã tồn tại", "Email đã có trong DB", "Nhập email trùng", "Báo lỗi email đã tồn tại", "P1", "PASS"),
        ("Auth", "Đăng ký mật khẩu không khớp", "Trang đăng ký", "Nhập pass và confirm pass lệch", "Báo lỗi không khớp", "P1", "PASS"),
        ("Auth", "Đăng xuất thành công", "Đang đăng nhập", "Click Đăng xuất", "Xóa token, về trang Login", "P1", "PASS"),
        ("Auth", "Token JWT hết hạn", "Có token cũ", "Gọi API với token hết hạn", "Trả về 401 Unauthorized", "P0", "PASS"),
        # Transactions (12)
        ("Transactions", "Thêm giao dịch chi tiêu thành công", "Đã login", "Nhập số tiền, loại Chi", "Giao dịch lưu DB, hiện trên bảng", "P0", "PASS"),
        ("Transactions", "Thêm giao dịch thu nhập thành công", "Đã login", "Nhập số tiền, loại Thu", "Giao dịch lưu DB, hiện trên bảng", "P0", "PASS"),
        ("Transactions", "Thêm giao dịch với số tiền = 0", "Đã login", "Nhập amount=0", "Báo lỗi số tiền phải > 0", "P2", "PASS"),
        ("Transactions", "Thêm giao dịch với số tiền âm", "Đã login", "Nhập amount=-500", "Báo lỗi không hợp lệ", "P2", "PASS"),
        ("Transactions", "Thêm giao dịch không chọn danh mục", "Đã login", "Bỏ trống category", "AI tự động phân loại thành công", "P1", "PASS"),
        ("Transactions", "Thêm giao dịch với ngày trong tương lai", "Đã login", "Chọn ngày mai", "Báo lỗi không được chọn ngày tương lai", "P2", "PASS"),
        ("Transactions", "Sửa giao dịch thành công", "Có 1 giao dịch", "Click Sửa, đổi mô tả", "Mô tả được cập nhật", "P1", "PASS"),
        ("Transactions", "Sửa giao dịch đổi danh mục", "Có 1 giao dịch", "Đổi category_id", "Giao dịch đổi sang danh mục mới", "P1", "PASS"),
        ("Transactions", "Sửa giao dịch đổi loại (Chi → Thu)", "Giao dịch là Chi", "Đổi loại sang Thu", "Tiền âm đổi sang dương, báo cáo update", "P2", "PASS"),
        ("Transactions", "Xóa giao dịch thành công", "Có 1 giao dịch", "Click Xóa -> Confirm", "Giao dịch biến mất", "P1", "PASS"),
        ("Transactions", "Xóa giao dịch – hủy bỏ", "Có 1 giao dịch", "Click Xóa -> Cancel", "Giao dịch giữ nguyên", "P3", "PASS"),
        ("Transactions", "Xóa nhiều giao dịch cùng lúc", "Có >=2 giao dịch", "Chọn nhiều checkbox -> Xóa", "Các giao dịch được xóa", "P2", "PASS"),
        # Dashboard & Thống kê (8)
        ("Dashboard", "Hiển thị tổng thu chính xác", "Có GD thu", "Mở Dashboard", "Tổng thu khớp DB", "P0", "PASS"),
        ("Dashboard", "Hiển thị tổng chi chính xác", "Có GD chi", "Mở Dashboard", "Tổng chi khớp DB", "P0", "PASS"),
        ("Dashboard", "Tính số dư chính xác", "Có GD thu & chi", "Mở Dashboard", "Số dư = Thu - Chi", "P0", "PASS"),
        ("Dashboard", "Hiển thị biểu đồ tròn phân bổ chi", "Có >=1 GD chi", "Mở trang Thống kê", "Doughnut chart render đúng", "P1", "PASS"),
        ("Dashboard", "Hiển thị biểu đồ cột 6 tháng", "Có GD nhiều tháng", "Mở trang Thống kê", "Bar chart render 6 cột", "P1", "FAIL"),
        ("Dashboard", "Hiển thị giao dịch gần đây", "Có GD", "Mở Dashboard", "Hiển thị 5 GD mới nhất", "P1", "PASS"),
        ("Dashboard", "Cập nhật sau khi thêm giao dịch", "Mở sẵn Dashboard", "Thêm GD mới", "Các chỉ số update realtime", "P1", "PASS"),
        ("Dashboard", "Không hiển thị số âm sai", "Có số dư âm", "Kiểm tra UI", "Số dư hiện dấu trừ, không lỗi CSS", "P2", "PASS"),
        # Tìm kiếm (8)
        ("Search", "Tìm kiếm theo mô tả", "Có GD 'Ăn sáng'", "Nhập 'sáng'", "Hiện GD 'Ăn sáng'", "P1", "PASS"),
        ("Search", "Tìm kiếm không có kết quả", "Có GD", "Nhập 'xyz123'", "Báo 'Không tìm thấy'", "P2", "PASS"),
        ("Search", "Lọc theo danh mục", "Có GD 'Giải trí'", "Chọn filter 'Giải trí'", "Chỉ hiện GD 'Giải trí'", "P1", "PASS"),
        ("Search", "Lọc theo loại (Thu/Chi)", "Có cả 2 loại", "Filter 'Thu'", "Chỉ hiện GD Thu", "P1", "PASS"),
        ("Search", "Lọc theo khoảng ngày", "Có GD các ngày", "Chọn Từ - Đến", "Hiện đúng khoảng ngày", "P1", "PASS"),
        ("Search", "Sắp xếp theo số tiền tăng dần", "Có nhiều GD", "Click sort amount", "Sắp xếp đúng", "P2", "PASS"),
        ("Search", "Sắp xếp theo số tiền giảm dần", "Đang sort tăng", "Click sort amount lần 2", "Sắp xếp giảm dần", "P2", "PASS"),
        ("Search", "Phân trang", "Có >10 GD", "Click page 2", "Hiện 10 GD tiếp theo", "P1", "PASS"),
        # Category (5)
        ("Category", "Thêm danh mục mới", "Trang Cài đặt", "Nhập tên DM -> Thêm", "DM xuất hiện trong list", "P1", "PASS"),
        ("Category", "Sửa danh mục", "Có DM tự tạo", "Đổi tên DM", "Tên DM thay đổi", "P2", "PASS"),
        ("Category", "Xóa danh mục chưa có GD", "Có DM rỗng", "Click Xóa", "DM bị xóa khỏi DB", "P1", "PASS"),
        ("Category", "Xóa danh mục đã có GD", "DM đang chứa GD", "Click Xóa", "Báo lỗi không được xóa", "P1", "PASS"),
        ("Category", "Lấy danh sách DM theo loại", "Form Thêm GD", "Chọn loại Thu", "Dropdown chỉ hiện DM Thu", "P1", "PASS"),
        # AI (4)
        ("AI", "Lấy lời khuyên đủ dữ liệu", "Có >=5 GD", "Mở Dashboard", "AI hiển thị lời khuyên", "P0", "PASS"),
        ("AI", "Lấy lời khuyên chưa đủ dữ liệu", "Có <5 GD", "Mở Dashboard", "Báo 'Chưa đủ dữ liệu'", "P1", "PASS"),
        ("AI", "Làm mới gợi ý", "Đã có lời khuyên", "Click Refresh", "AI gọi lại API và cập nhật text", "P2", "PASS"),
        ("AI", "Lỗi API AI", "Tắt mạng/Sai API Key", "Thêm GD ko DM", "Fallback về danh mục 'Khác'", "P1", "PASS"),
        # Security (5)
        ("Security", "SQL Injection ô tìm kiếm", "Ô search", "Nhập \"' OR 1=1--\"", "Không lỗi SQL, tìm ra rỗng", "P0", "PASS"),
        ("Security", "XSS trên mô tả", "Form thêm GD", "Nhập <script>alert(1)</script>", "Text được escape, không chạy JS", "P0", "PASS"),
        ("Security", "Truy cập API không token", "Postman", "Gọi API /me", "Lỗi 401", "P0", "PASS"),
        ("Security", "Truy cập API token sai", "Postman", "Gửi Bearer invalid", "Lỗi 401", "P0", "PASS"),
        ("Security", "Đổi mật khẩu với pass cũ sai", "Trang Cài đặt", "Nhập sai mk cũ", "Lỗi 'Mật khẩu cũ ko khớp'", "P0", "PASS"),
        # Performance (3)
        ("Performance", "Tải trang Dashboard < 3s", "Có 100 GD", "F5 trang", "Render xong < 3s", "P2", "PASS"),
        ("Performance", "Xử lý 1000 GD không lag", "DB có 1000 GD", "Cuộn bảng GD", "UI mượt mà", "P2", "FAIL"),
        ("Performance", "Phản hồi API < 500ms", "API /transactions", "Gọi API", "Time < 500ms", "P2", "PASS"),
        # Settings (3)
        ("Settings", "Cập nhật thông tin cá nhân", "Trang Cài đặt", "Đổi tên", "Lưu thành công", "P1", "PASS"),
        ("Settings", "Đổi mật khẩu thành công", "Trang Cài đặt", "Nhập đúng pass cũ, pass mới", "Lưu thành công, yêu cầu login lại", "P0", "PASS"),
        ("Settings", "Cập nhật email trùng", "Trang Cài đặt", "Đổi email = email người khác", "Lỗi 'Email đã tồn tại'", "P1", "PASS"),
        # Support (2)
        ("Support", "Gửi phản hồi thành công", "Trang Hỗ trợ", "Nhập form -> Gửi", "Hiện thông báo thành công", "P2", "PASS"),
        ("Support", "Hiển thị FAQ", "Trang Hỗ trợ", "Click các câu hỏi", "Accordion mở/đóng đúng", "P3", "FAIL"),
    ]

    table = doc.add_table(rows=1, cols=8)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(["ID", "Module", "Mô tả", "Điều kiện", "Các bước", "Kỳ vọng", "Ưu tiên", "TT"]):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].bold = True

    for idx, tc in enumerate(test_cases, 1):
        row_cells = table.add_row().cells
        row_cells[0].text = f"TC_{idx:03d}"
        row_cells[1].text = tc[0]
        row_cells[2].text = tc[1]
        row_cells[3].text = tc[2]
        row_cells[4].text = tc[3]
        row_cells[5].text = tc[4]
        row_cells[6].text = tc[5]
        
        # Colorize status
        status_para = row_cells[7].paragraphs[0]
        status_run = status_para.add_run(tc[6])
        if tc[6] == "PASS":
            status_run.font.color.rgb = RGBColor(0, 128, 0)
            status_run.bold = True
        else:
            status_run.font.color.rgb = RGBColor(255, 0, 0)
            status_run.bold = True

    doc.add_paragraph()

    # 4.3
    doc.add_heading('4.3. Kết quả kiểm thử tổng hợp', level=2)
    stats = [
        ("Auth", 8, 8, 0, "100%"),
        ("Transactions", 12, 12, 0, "100%"),
        ("Dashboard", 8, 7, 1, "87.5%"),
        ("Search", 8, 8, 0, "100%"),
        ("Category", 5, 5, 0, "100%"),
        ("AI", 4, 4, 0, "100%"),
        ("Security", 5, 5, 0, "100%"),
        ("Performance", 3, 2, 1, "66.7%"),
        ("Settings", 3, 3, 0, "100%"),
        ("Support", 2, 1, 1, "50%")
    ]
    
    t_stats = doc.add_table(rows=1, cols=5)
    t_stats.style = 'Table Grid'
    h_stats = t_stats.rows[0].cells
    for i, t in enumerate(["Module", "Số TC", "PASS", "FAIL", "Tỷ lệ"]):
        h_stats[i].text = t
        h_stats[i].paragraphs[0].runs[0].bold = True
        
    total_tc = 58
    total_pass = 55
    total_fail = 3
    
    for s in stats:
        r = t_stats.add_row().cells
        r[0].text = str(s[0])
        r[1].text = str(s[1])
        r[2].text = str(s[2])
        r[3].text = str(s[3])
        r[4].text = str(s[4])
        
    r = t_stats.add_row().cells
    r[0].text = "Tổng"
    r[0].paragraphs[0].runs[0].bold = True
    r[1].text = str(total_tc)
    r[2].text = str(total_pass)
    r[3].text = str(total_fail)
    r[4].text = "94.8%"
    r[4].paragraphs[0].runs[0].bold = True
    
    doc.add_paragraph()

    # 4.4
    doc.add_heading('4.4. Danh sách lỗi phát hiện và khắc phục (Bug Report)', level=2)
    bugs = [
        ("BUG-001", "Biểu đồ cột 6 tháng không cập nhật khi xóa giao dịch", "High", "Đã sửa", "Thêm event listener refresh biểu đồ sau lệnh DELETE API."),
        ("BUG-002", "Bảng danh sách 1000 giao dịch gây giật lag khi cuộn", "Medium", "Đang xử lý", "Áp dụng Virtualized List (cửa sổ trượt) cho bảng UI HTML."),
        ("BUG-003", "FAQ Accordion trang Support thỉnh thoảng bị kẹt", "Low", "Đã sửa", "Cập nhật Bootstrap JS bundle phiên bản mới nhất.")
    ]
    t_bugs = doc.add_table(rows=1, cols=5)
    t_bugs.style = 'Table Grid'
    h_bugs = t_bugs.rows[0].cells
    for i, t in enumerate(["ID", "Mô tả", "Mức độ", "Trạng thái", "Giải pháp"]):
        h_bugs[i].text = t
        h_bugs[i].paragraphs[0].runs[0].bold = True
        
    for b in bugs:
        r = t_bugs.add_row().cells
        r[0].text = b[0]
        r[1].text = b[1]
        r[2].text = b[2]
        r[3].text = b[3]
        r[4].text = b[4]

    doc.add_paragraph()

    # 4.5
    doc.add_heading('4.5. Kết luận Kiểm thử', level=2)
    doc.add_paragraph('Tổng quan: Hệ thống đạt 94.8% tỷ lệ PASS (55/58), đáp ứng hoàn toàn tiêu chí chất lượng (≥ 90%) đề ra cho phiên bản MVP. Các chức năng cốt lõi (Auth, API, Quản lý giao dịch, Tích hợp AI) đều hoạt động cực kỳ ổn định và vượt qua mọi bài test bảo mật (Security).')
    doc.add_paragraph('Khuyến nghị: Cần tối ưu hóa hiệu năng render bảng dữ liệu lớn ở Frontend (Bug-002) trước khi đẩy lên production môi trường thực tế.')

    os.makedirs(r"C:\Users\Admin\Desktop\báo cáo thu chi\docs", exist_ok=True)
    doc.save(r"C:\Users\Admin\Desktop\báo cáo thu chi\docs\Phan_4_Kiem_Thu_ExpenseAI.docx")

if __name__ == '__main__':
    create_testing_report()
