"""
Script thay thế đoạn ASCII text wireframe bị lỗi vỡ trang trong Bao_Cao_3_Bai_Kiem_Tra_Backup.docx
bằng một Bảng Wireframe thiết kế chuẩn mực, chuyên nghiệp và khóa ngắt trang (cantSplit).
"""
import sys
import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    cell._tc.get_or_add_tcPr().append(tcMar)

def set_table_borders(table, color_hex="94A3B8"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/><w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color_hex}"/><w:left w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/><w:right w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/><w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/></w:tblBorders>')
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def fix_wireframe(docx_path):
    d = docx.Document(docx_path)
    
    target_idx = None
    for i, p in enumerate(d.paragraphs):
        if "3.4.4. Wireframe/Prototype" in p.text:
            target_idx = i + 1
            break
            
    if target_idx is None or target_idx >= len(d.paragraphs):
        print("Không tìm thấy vị trí 3.4.4 trong tài liệu!")
        return False
        
    p_target = d.paragraphs[target_idx]
    # Sửa lại nội dung tiêu đề đoạn văn
    p_target.text = "Mô tả cấu trúc khung lưới (Wireframe Layout) của trang chủ Dashboard:"
    p_target.paragraph_format.space_before = Pt(4)
    p_target.paragraph_format.space_after = Pt(8)
    p_target.runs[0].font.name = "Times New Roman"
    p_target.runs[0].font.size = Pt(13)
    p_target.runs[0].font.italic = True
    
    # Tạo bảng mô phỏng Wireframe chuyên nghiệp
    # Để chèn ngay sau p_target, ta tạo bảng rồi di chuyển node XML của bảng vào sau p_target._p
    table = d.add_table(rows=4, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, "CBD5E1")
    
    # Di chuyển table vào đúng vị trí sau p_target
    p_target._p.addnext(table._tbl)
    
    # Cấu hình độ rộng các cột
    widths = [Cm(5.0), Cm(5.0), Cm(5.5)]
    for row in table.rows:
        make_row_cant_split(row)
        for i, w in enumerate(widths):
            row.cells[i].width = w
            
    # Hàng 0: Header Navbar (Merge cả 3 cột)
    c_nav = table.cell(0, 0)
    c_nav.merge(table.cell(0, 1)).merge(table.cell(0, 2))
    set_cell_background(c_nav, "1E293B")
    set_cell_margins(c_nav, 160, 160, 180, 180)
    p_nav = c_nav.paragraphs[0]
    p_nav.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_nav = p_nav.add_run("[ ExpenseAI ]       Dashboard       Sổ Giao Dịch       Danh Mục       [ User: Tuấn Đạt ]")
    r_nav.font.name = "Times New Roman"
    r_nav.font.size = Pt(11)
    r_nav.font.bold = True
    r_nav.font.color.rgb = RGBColor(255, 255, 255)
    
    # Hàng 1: 3 Thẻ chỉ số tài chính (Cards)
    cards_data = [
        ("THẺ: TỔNG THU", "+15,000,000 đ", "ECFDF5", RGBColor(5, 150, 105)),
        ("THẺ: TỔNG CHI", "-8,450,000 đ", "FFF1F2", RGBColor(225, 29, 72)),
        ("THẺ: SỐ DƯ KHẢ DỤNG", "+6,550,000 đ", "EFF6FF", RGBColor(37, 99, 235)),
    ]
    for col_idx, (title, val, bg, text_color) in enumerate(cards_data):
        cell = table.cell(1, col_idx)
        set_cell_background(cell, bg)
        set_cell_margins(cell, 140, 140, 140, 140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rt = p.add_run(f"{title}\n")
        rt.font.name = "Times New Roman"
        rt.font.size = Pt(9.5)
        rt.font.bold = True
        rt.font.color.rgb = RGBColor(100, 116, 139)
        rv = p.add_run(val)
        rv.font.name = "Times New Roman"
        rv.font.size = Pt(13)
        rv.font.bold = True
        rv.font.color.rgb = text_color
        
    # Hàng 2: AI Prompt Input (Merge cả 3 cột)
    c_ai = table.cell(2, 0)
    c_ai.merge(table.cell(2, 1)).merge(table.cell(2, 2))
    set_cell_background(c_ai, "F8FAFC")
    set_cell_margins(c_ai, 140, 140, 180, 180)
    p_ai = c_ai.paragraphs[0]
    p_ai.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ai_lbl = p_ai.add_run("[ AI PROMPT INPUT ]: ")
    r_ai_lbl.font.name = "Times New Roman"
    r_ai_lbl.font.size = Pt(10.5)
    r_ai_lbl.font.bold = True
    r_ai_lbl.font.color.rgb = RGBColor(30, 41, 59)
    r_ai_val = p_ai.add_run('"Hôm nay đi ăn tối với bạn hết 120k"    ')
    r_ai_val.font.name = "Times New Roman"
    r_ai_val.font.size = Pt(10.5)
    r_ai_val.font.italic = True
    r_ai_val.font.color.rgb = RGBColor(71, 85, 105)
    r_ai_btn = p_ai.add_run("[ Ghi Nhận Ngay ]")
    r_ai_btn.font.name = "Times New Roman"
    r_ai_btn.font.size = Pt(10.5)
    r_ai_btn.font.bold = True
    r_ai_btn.font.color.rgb = RGBColor(26, 54, 93)

    # Hàng 3: Hai khu vực chức năng (Budget Progress & Chart Breakdown)
    c_bgt = table.cell(3, 0)
    c_bgt.merge(table.cell(3, 1))
    set_cell_background(c_bgt, "FFFFFF")
    set_cell_margins(c_bgt, 140, 140, 160, 160)
    p_bgt = c_bgt.paragraphs[0]
    r_bgt_t = p_bgt.add_run("TIẾN ĐỘ NGÂN SÁCH THÁNG (70.4%)\n")
    r_bgt_t.font.name = "Times New Roman"
    r_bgt_t.font.size = Pt(10.5)
    r_bgt_t.font.bold = True
    r_bgt_t.font.color.rgb = RGBColor(30, 41, 59)
    r_bgt_v = p_bgt.add_run("• Đã chi tiêu: 8.450.000 đ / 12.000.000 đ\n• Trạng thái: An toàn (Còn lại 3.550.000 đ)\n• Tỷ lệ đạt ngưỡng: 70.4% (Chưa vượt trần cảnh báo)")
    r_bgt_v.font.name = "Times New Roman"
    r_bgt_v.font.size = Pt(10)
    r_bgt_v.font.color.rgb = RGBColor(71, 85, 105)

    c_cht = table.cell(3, 2)
    set_cell_background(c_cht, "FFFFFF")
    set_cell_margins(c_cht, 140, 140, 160, 160)
    p_cht = c_cht.paragraphs[0]
    r_cht_t = p_cht.add_run("CƠ CẤU CHI TIÊU (Chart.js)\n")
    r_cht_t.font.name = "Times New Roman"
    r_cht_t.font.size = Pt(10.5)
    r_cht_t.font.bold = True
    r_cht_t.font.color.rgb = RGBColor(30, 41, 59)
    r_cht_v = p_cht.add_run("• Ăn uống: 45% (3.800.000 đ)\n• Tiền phòng trọ: 30% (2.500.000 đ)\n• Di chuyển, xăng xe: 15% (1.250.000 đ)\n• Chi tiêu khác: 10% (900.000 đ)")
    r_cht_v.font.name = "Times New Roman"
    r_cht_v.font.size = Pt(10)
    r_cht_v.font.color.rgb = RGBColor(71, 85, 105)

    # Chèn Caption hình ảnh ngay sau bảng
    p_caption = d.add_paragraph()
    p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_caption.paragraph_format.space_before = Pt(6)
    p_caption.paragraph_format.space_after = Pt(12)
    r_cap = p_caption.add_run("Hình 3.6: Bản phác thảo cấu trúc khung lưới (Wireframe) Dashboard hệ thống ExpenseAI")
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = RGBColor(100, 116, 139)
    
    # Di chuyển caption vào sau table
    table._tbl.addnext(p_caption._p)
    
    d.save(docx_path)
    print("Fix wireframe completed successfully!")
    return True

if __name__ == "__main__":
    fix_wireframe("c:/Users/dathao/Downloads/AI/Bao_Cao_3_Bai_Kiem_Tra_Backup_Da_Sua.docx")

