import docx

doc = docx.Document(r"c:\Users\dathao\Downloads\AI\Bao_Cao_3_Bai_Kiem_Tra_Backup_Da_Sua.docx")
for i in range(144, 155):
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        xml = p._element.xml
        has_draw = "w:drawing" in xml
        has_br = 'w:type="page"' in xml
        txt_safe = p.text[:65].encode('ascii', 'backslashreplace').decode('ascii')
        print(f"{i:3d} | Draw={str(has_draw):5s} | PageBreak={str(has_br):5s} | Text: {txt_safe}")
