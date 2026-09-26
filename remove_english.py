import docx
import re

def should_replace(match_text):
    inner = match_text.strip()[1:-1].strip()
    if inner.isupper():
        return False
    if inner in ["ICTU", "ERD", "OWASP", "RBAC", "OOP", "OOM", "SOLID", "API", "UI", "UX", "AI"]:
        return False
    return True

def replace_in_paragraph(p, regex):
    while True:
        text = p.text
        match = regex.search(text)
        if not match:
            break
            
        if not should_replace(match.group(0)):
            # If we don't want to replace, we need a way to skip it.
            # Easiest way is to temporarily replace the text with a placeholder, then revert later, 
            # but that's complex with runs.
            # Instead, we can search for all matches, filter them, and replace from right to left!
            break # This logic is flawed if we just break. 
        
        # We need right to left replacement
        pass

def process_paragraphs(paragraphs, regex):
    for p in paragraphs:
        while True:
            text = p.text
            matches = list(regex.finditer(text))
            valid_matches = [m for m in matches if should_replace(m.group(0))]
            
            if not valid_matches:
                break
                
            # Take the last valid match to avoid messing up indices for earlier matches
            match = valid_matches[-1]
            start, end = match.span()
            
            current_idx = 0
            for run in p.runs:
                run_len = len(run.text)
                run_start = current_idx
                run_end = current_idx + run_len
                
                if run_end > start and run_start < end:
                    keep_start = ""
                    keep_end = ""
                    
                    if run_start < start:
                        keep_start = run.text[:start - run_start]
                    if run_end > end:
                        keep_end = run.text[end - run_start:]
                        
                    run.text = keep_start + keep_end
                    
                current_idx += run_len

def process_doc():
    doc = docx.Document('docs/Bao_Cao_Hoan_Chinh_ExpenseAI.docx')
    regex = re.compile(r' \([A-Za-z\s-]+\)')
    
    process_paragraphs(doc.paragraphs, regex)
    
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                process_paragraphs(cell.paragraphs, regex)
                
    doc.save('docs/Bao_Cao_Hoan_Chinh_ExpenseAI.docx')
    print("Done removing English translations.")

if __name__ == '__main__':
    process_doc()
