import urllib.request
import base64
import zlib
import os
import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# The 4 Mermaid diagrams generated in the previous step
diagrams = {
    'UseCase': '''flowchart LR
    classDef actor fill:#F5F3FF,stroke:#7C3AED,stroke-width:2px,color:#0F172A
    classDef usecase fill:#F0F7FF,stroke:#1F4E79,stroke-width:2px,color:#0F172A,rx:20,ry:20
    classDef system fill:#FAFAFA,stroke:#E2E8F0,stroke-width:2px,stroke-dasharray: 5 5

    User((Người dùng)):::actor
    AI((Gemini AI)):::actor

    subgraph System [ExpenseAI Web Application]
        direction TB
        UC1(Đăng ký & Đăng nhập):::usecase
        UC2(Thêm/Sửa/Xóa Giao dịch):::usecase
        UC3(Xem Báo cáo & Dashboard):::usecase
        UC4(Phân loại Giao dịch Tự động):::usecase
        UC5(Tư vấn Tài chính Cá nhân):::usecase
    end

    User --> UC1
    User --> UC2
    User --> UC3
    User --> UC5

    UC2 -. "<< include >>" .-> UC4
    UC4 <--> AI
    UC5 <--> AI
    class System system''',

    'ERD': '''erDiagram
    USERS {
        UUID id PK
        VARCHAR email
        VARCHAR password_hash
        TIMESTAMP created_at
    }
    TRANSACTIONS {
        UUID id PK
        UUID user_id FK
        DECIMAL amount
        VARCHAR type
        VARCHAR description
        VARCHAR ai_category
        TIMESTAMP date
    }
    AI_ADVICE {
        UUID id PK
        UUID user_id FK
        TEXT advice_content
        TIMESTAMP generated_at
    }
    USERS ||--o{ TRANSACTIONS : "tạo"
    USERS ||--o{ AI_ADVICE : "nhận"''',

    'Architecture': '''flowchart TD
    classDef client fill:#F8FAFC,stroke:#334155,stroke-width:2px,color:#0F172A
    classDef server fill:#F0F7FF,stroke:#2563EB,stroke-width:2px,color:#0F172A
    classDef db fill:#ECFDF5,stroke:#059669,stroke-width:2px,color:#0F172A
    classDef ai fill:#F5F3FF,stroke:#7C3AED,stroke-width:2px,color:#0F172A

    subgraph Tier1 [Presentation Tier - Frontend]
        UI[Trình duyệt Web\nHTML/CSS/JS/Bootstrap 5]:::client
    end

    subgraph Tier2 [Application Tier - Backend]
        API[FastAPI Router\nJWT Auth]:::server
        Service[AI Logic & Business]:::server
        ORM[SQLAlchemy ORM]:::server
        API --> Service --> ORM
    end

    subgraph Tier3 [Data Tier]
        DB[(PostgreSQL 15+)]:::db
    end
    
    subgraph External [External Services]
        LLM{Gemini 1.5 Flash\nOpenAI Fallback}:::ai
    end

    UI <-->|HTTP/REST API| API
    ORM <-->|TCP/IP: 5432| DB
    Service <-->|HTTPS API Call| LLM''',

    'Activity': '''stateDiagram-v2
    state "Người dùng nhập Giao dịch" as Input
    state "FastAPI tiếp nhận" as API
    state "Gọi Gemini AI" as Gemini
    state "AI Phân loại danh mục" as Classify
    state "Lưu vào PostgreSQL" as Save
    state "Cập nhật Dashboard" as Dashboard
    
    [*] --> Input
    Input --> API : Submit Form
    API --> Gemini : Gửi mô tả
    Gemini --> Classify : Trả về JSON
    Classify --> Save : Map với User ID
    Save --> Dashboard : Trả về kết quả
    Dashboard --> [*]'''
}

img_dir = r'C:\Users\dathao\Desktop\ảnh mới vẽ'
os.makedirs(img_dir, exist_ok=True)

print("Starting image generation...")
file_paths = {}
for name, code in diagrams.items():
    try:
        compressed = zlib.compress(code.encode('utf-8'))
        b64 = base64.urlsafe_b64encode(compressed).decode('utf-8')
        url = 'https://kroki.io/mermaid/png/' + b64
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            file_path = os.path.join(img_dir, f'new_{name}.png')
            with open(file_path, 'wb') as f:
                f.write(response.read())
            file_paths[name] = file_path
        print(f"Generated {name}.png")
    except Exception as e:
        print(f"Failed to generate {name}: {e}")

doc_path = r'C:\Users\dathao\Desktop\nhom02_FINAL_backup_before_videofix.docx'
out_path = r'C:\Users\dathao\Desktop\nhom02_FINAL_updated_with_diagrams.docx'
try:
    doc = docx.Document(doc_path)
    print("Document loaded.")
    
    for p in doc.paragraphs:
        text = p.text.strip().lower()
        
        # Match captions to insert our generated images
        if text.startswith('hình 2.1a: biểu đồ use case'):
            if 'UseCase' in file_paths:
                new_p = p.insert_paragraph_before()
                new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = new_p.add_run()
                run.add_picture(file_paths['UseCase'], width=Inches(6.0))
                print("Inserted UseCase before Hình 2.1a")
                
        elif text.startswith('hình 2.6: sơ đồ thực thể'):
            if 'ERD' in file_paths:
                new_p = p.insert_paragraph_before()
                new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = new_p.add_run()
                run.add_picture(file_paths['ERD'], width=Inches(6.0))
                print("Inserted ERD before Hình 2.6")
                
        elif text.startswith('hình 3.1: sơ đồ kiến trúc'):
            if 'Architecture' in file_paths:
                new_p = p.insert_paragraph_before()
                new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = new_p.add_run()
                run.add_picture(file_paths['Architecture'], width=Inches(6.0))
                print("Inserted Architecture before Hình 3.1")
                
        elif text.startswith('hình 2.2a: sơ đồ hoạt động'):
            if 'Activity' in file_paths:
                new_p = p.insert_paragraph_before()
                new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = new_p.add_run()
                run.add_picture(file_paths['Activity'], width=Inches(6.0))
                print("Inserted Activity before Hình 2.2a")
                
    doc.save(out_path)
    print("Done! Saved to", out_path)
except Exception as e:
    print('Error:', e)
