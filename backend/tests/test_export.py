"""Unit tests cho tính năng Xuất báo cáo Excel (Export)."""
from datetime import timedelta
import io
import openpyxl
import pytest
from fastapi.testclient import TestClient

from src.database import SessionLocal
from src.main import app
from src.models import User, Role
from src.utils.security import create_access_token, get_password_hash

client = TestClient(app)


@pytest.fixture
def user_token():
    db = SessionLocal()
    user = db.query(User).filter(User.username != "admin").first()
    if not user:
        user = User(
            username="export_user_test",
            email="export_user@example.com",
            hashed_password=get_password_hash("password123"),
        )
        user_role = db.query(Role).filter(Role.name == "user").first()
        if user_role:
            user.roles.append(user_role)
        db.add(user)
        db.commit()
        db.refresh(user)
    token = create_access_token({"sub": str(user.id)}, expires_delta=timedelta(minutes=30))
    db.close()
    return token


def test_export_excel_success(user_token):
    """Kiểm thử xuất báo cáo Excel thành công, trả về đúng định dạng MIME và workbook hợp lệ."""
    headers = {"Authorization": f"Bearer {user_token}"}
    res = client.get("/api/reports/export/excel?month=2026-09", headers=headers)
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert "ExpenseAI_Report" in res.headers["content-disposition"]

    # Đọc nội dung workbook từ bytes trả về
    excel_stream = io.BytesIO(res.content)
    wb = openpyxl.load_workbook(excel_stream)
    assert "Giao Dịch" in wb.sheetnames
    ws = wb["Giao Dịch"]
    # Kiểm tra header và dữ liệu
    assert "BÁO CÁO CHI TIẾT GIAO DỊCH EXPENSEAI" in str(ws["A1"].value)


def test_export_excel_unauthenticated():
    """Từ chối yêu cầu xuất báo cáo khi chưa xác thực."""
    res = client.get("/api/reports/export/excel")
    assert res.status_code == 401
