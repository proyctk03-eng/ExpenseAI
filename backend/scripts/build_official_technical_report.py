from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Bao_Cao_Ky_Thuat_Chinh_Thuc.docx"


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), "D9D9D9")
        borders.append(tag)
    tc_pr.append(borders)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        set_cell_fill(cell, "17365D")
        set_cell_border(cell)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            cell.text = str(value)
            if row_index % 2:
                set_cell_fill(cell, "F4F7FA")
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            if widths:
                cell.width = Cm(widths[idx])
    doc.add_paragraph()


def bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "Arial"
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Báo cáo kỹ thuật chính thức ExpenseAI")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Phiên bản đã đối chiếu mã nguồn và kiểm thử - 17/09/2026").italic = True

    doc.add_heading("Kết luận", level=1)
    doc.add_paragraph(
        "ExpenseAI là ứng dụng web quản lý thu chi cá nhân có tích hợp AI. Báo cáo này thay thế các mô tả không còn khớp với mã nguồn. "
        "Nó chỉ ghi nhận chức năng, kiến trúc và kết quả kiểm thử đã được kiểm chứng tại thời điểm phát hành."
    )

    doc.add_heading("Phạm vi và kiến trúc", level=1)
    doc.add_paragraph(
        "Hệ thống hỗ trợ xác thực JWT, CRUD giao dịch và danh mục, dashboard, báo cáo, quét hóa đơn, tư vấn AI, phân tích hành vi có xác nhận, feedback ticket và bộ nhớ danh mục theo người dùng. "
        "Quản lý ngân sách và cảnh báo ngưỡng chưa được triển khai."
    )
    add_table(doc, ["Lớp", "Thành phần đã triển khai"], [
        ["Giao diện", "Jinja2, Bootstrap, Chart.js"],
        ["Ứng dụng", "FastAPI routers, Pydantic, phân quyền RBAC"],
        ["Dịch vụ", "Classifier, advice, behavior, vision; fallback nội bộ"],
        ["Dữ liệu", "SQLAlchemy ORM; PostgreSQL hoặc SQLite; Redis tùy chọn"],
        ["AI", "Gemini 1.5 Flash ưu tiên; OpenAI GPT-3.5 Turbo fallback"],
    ], [3.5, 12.5])

    doc.add_heading("Dữ liệu và API", level=1)
    doc.add_paragraph(
        "Schema hiện hành có 11 bảng: users, categories, transactions, ai_predictions, user_memory_rules, roles, permissions, user_roles, role_permissions, feedback_tickets và ticket_replies. "
        "User và role có quan hệ nhiều-nhiều; giao dịch thuộc user và tùy chọn category; AI prediction là audit một-một với giao dịch."
    )
    add_table(doc, ["Nhóm", "Endpoint chính"], [
        ["Auth", "POST /api/auth/register, /login, /refresh, /logout; GET/PUT /api/auth/me"],
        ["Giao dịch", "GET/POST /api/transactions/; PUT/DELETE /api/transactions/{id}; POST /scan-receipt"],
        ["Báo cáo", "GET /api/reports/summary, /by_category, /monthly_trend; GET /api/dashboard/*"],
        ["AI", "POST /api/advice/; GET /api/advice/behavior?share_transaction_details=true"],
        ["Feedback", "Các endpoint tại /api/feedback/ cho ticket, trạng thái và phản hồi"],
    ], [3.5, 12.5])

    doc.add_heading("Quyền riêng tư và bảo mật", level=1)
    bullet(doc, "Lời khuyên tài chính mặc định chỉ gửi số liệu đã tổng hợp theo danh mục.")
    bullet(doc, "Phân tích hành vi có thể gửi mô tả, ngày và số tiền của tối đa 50 giao dịch chi. Người dùng phải xác nhận rõ ràng trước khi gọi.")
    bullet(doc, "Access token và refresh token được đặt trong HttpOnly cookie; API vẫn hỗ trợ Bearer token cho client API.")
    bullet(doc, "Cấu hình production yêu cầu SECRET_KEY, DATABASE_URL và ít nhất một khóa Gemini hoặc OpenAI.")

    doc.add_heading("Trạng thái kiểm thử và triển khai", level=1)
    doc.add_paragraph(
        "Lệnh kiểm thử đã chạy: OPENAI_API_KEY=mock-api-key-for-testing, GEMINI_API_KEY=mock-api-key-for-testing, CACHE_ENABLED=false, python -m pytest -q. "
        "Kết quả: 50 passed. Cấu hình giả lập giúp kiểm thử không gọi dịch vụ AI bên ngoài."
    )
    bullet(doc, "CI workflow tại .github/workflows/ci.yml chạy test trên Python 3.10, 3.11 và 3.12, đồng thời quét Trivy. Chưa khẳng định trạng thái CI thành công trước khi workflow chạy trên GitHub.")
    bullet(doc, "Alembic có migration tạo schema đầy đủ. Docker Compose dùng service backend và Nginx đã trỏ đúng upstream này.")
    bullet(doc, "Không công bố các chỉ số 96%, 100%, zero downtime hoặc ngưỡng hiệu năng nếu chưa có benchmark tái lập được.")

    doc.add_heading("Quản trị tài liệu", level=1)
    doc.add_paragraph(
        "Nguồn chuẩn ở docs/00_SOURCE_OF_TRUTH.md. Tài liệu SDLC được cập nhật theo endpoint và schema thực tế. Các tệp trong docs/backup/ và docs/backup/archive_old/ là lịch sử, không phải tài liệu nộp."
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    main()
