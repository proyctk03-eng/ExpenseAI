import docx
doc_path = r'C:\Users\dathao\Desktop\Bao_Cao_Hoan_Chinh_ExpenseAI.docx'
try:
    doc = docx.Document(doc_path)
    with open('placeholders.txt', 'w', encoding='utf-8') as f:
        count = 0
        for i, p in enumerate(doc.paragraphs):
            if '[CHÈN SƠ ĐỒ MỚI TẠI ĐÂY]' in p.text:
                count += 1
                f.write(f"Placeholder {count}:\n")
                for j in range(max(0, i-2), min(len(doc.paragraphs), i+3)):
                    if doc.paragraphs[j].text.strip():
                        f.write(f"  {doc.paragraphs[j].text.strip()}\n")
                f.write("-" * 20 + "\n")
except Exception as e:
    print('Error:', e)
