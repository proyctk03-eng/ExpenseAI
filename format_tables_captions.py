import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import time

doc_path = r'C:\Users\dathao\Desktop\nhom02_FINAL_backup_before_videofix.docx'

try:
    doc = docx.Document(doc_path)
    print("Loaded document successfully.")
    
    # 1. FORMAT TABLES
    for table in doc.tables:
        table.style = 'Table Grid' # Applies standard solid borders
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    if i == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        
                    for run in p.runs:
                        run.font.name = 'Times New Roman'
                        rPr = run._element.get_or_add_rPr()
                        rFonts = rPr.get_or_add_rFonts()
                        rFonts.set(qn('w:eastAsia'), 'Times New Roman')
                        
                        if i == 0:
                            run.bold = True
                        else:
                            run.bold = False

    # 2. MOVE CAPTIONS TO BOTTOM
    body_elements = doc._body._body
    
    # A. Move captions above tables to below tables
    moves_tbl = []
    for i in range(len(body_elements) - 1):
        el = body_elements[i]
        next_el = body_elements[i+1]
        
        if el.tag.endswith('p') and next_el.tag.endswith('tbl'):
            p = docx.text.paragraph.Paragraph(el, doc)
            text = p.text.strip().lower()
            if text.startswith('bảng ') or text.startswith('hình ') or text.startswith('biểu đồ'):
                moves_tbl.append((el, next_el))
                print(f"Found caption above table: {text[:30]}...")

    for p_el, tbl_el in moves_tbl:
        tbl_el.addnext(p_el)

    # B. Move captions above images to below images
    moves_img = []
    # Re-fetch elements since order changed
    body_elements = doc._body._body 
    for i in range(len(body_elements) - 1):
        el = body_elements[i]
        next_el = body_elements[i+1]
        
        if el.tag.endswith('p') and next_el.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(el, doc)
            p_next = docx.text.paragraph.Paragraph(next_el, doc)
            
            text = p.text.strip().lower()
            if text.startswith('bảng ') or text.startswith('hình ') or text.startswith('biểu đồ'):
                if '<w:drawing' in p_next._element.xml or '<v:imagedata' in p_next._element.xml:
                    moves_img.append((el, next_el))
                    print(f"Found caption above image: {text[:30]}...")

    for p_el, img_p_el in moves_img:
        img_p_el.addnext(p_el)

    # 3. FORMAT CAPTIONS
    for p in doc.paragraphs:
        text = p.text.strip().lower()
        if text.startswith('bảng ') or text.startswith('hình ') or text.startswith('biểu đồ '):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.italic = True
                
    doc.save(doc_path)
    print("Done formatting tables and moving captions.")

except Exception as e:
    print("Error:", e)
