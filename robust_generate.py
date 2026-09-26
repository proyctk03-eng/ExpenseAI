import urllib.request
import base64
import zlib
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

diagrams = {
    '2.1a': '''flowchart LR
    User((Người Dùng))
    Gemini((Google Gemini AI))
    User --> UC1([Đăng ký/Đăng nhập])
    User --> UC2([Quản lý Giao dịch])
    User --> UC3([Thiết lập Ngân sách])
    User --> UC4([Nhận Tư vấn Tài chính])
    UC2 -.->|Gọi API| Gemini
    UC4 -.->|Phân tích| Gemini'''
}

img_dir = 'C:/Users/dathao/Downloads/AI/ExpenseAI/scratch/images'
os.makedirs(img_dir, exist_ok=True)

print("Starting image generation...")
for name, code in diagrams.items():
    try:
        compressed = zlib.compress(code.encode('utf-8'))
        b64 = base64.urlsafe_b64encode(compressed).decode('utf-8')
        url = 'https://kroki.io/mermaid/png/' + b64
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=5) as response:
            with open(os.path.join(img_dir, f'Hình {name}.png'), 'wb') as f:
                f.write(response.read())
        print(f"Generated {name}")
    except Exception as e:
        print(f"Failed to generate {name}: {e}")

doc_path = r'C:\Users\dathao\Desktop\Bao_Cao_Hoan_Chinh_ExpenseAI.docx'
out_path = r'C:\Users\dathao\Desktop\Bao_Cao_Final_ExpenseAI.docx'
try:
    doc = docx.Document(doc_path)
    print("Document loaded.")
    
    # Insert cover page
    if doc.paragraphs[0].text != 'ĐẠI HỌC THÁI NGUYÊN':
        for section in doc.sections:
            section.top_margin = Inches(2.0 / 2.54)
            section.bottom_margin = Inches(2.0 / 2.54)
            section.left_margin = Inches(3.5 / 2.54)
            section.right_margin = Inches(2.0 / 2.54)
            
        first_p = doc.paragraphs[0]
        
        def add_para(text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, size=13, space_after=0, left_indent=0):
            new_p = first_p.insert_paragraph_before('')
            new_p.alignment = align
            new_p.paragraph_format.space_after = Pt(space_after)
            if left_indent > 0:
                new_p.paragraph_format.left_indent = Inches(left_indent / 2.54)
            run = new_p.add_run(text)
            run.bold = bold
            run.font.size = Pt(size)
            run.font.name = 'Times New Roman'
            run.font.color.rgb = RGBColor(0, 0, 0)
            return new_p

        add_para("ĐẠI HỌC THÁI NGUYÊN", bold=False, size=13)
        add_para("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG", bold=True, size=13)
        add_para("------------------***------------------", bold=False, size=11, space_after=36)
        add_para("BÁO CÁO HỌC PHẦN / BÀI TẬP LỚN", bold=True, size=16, space_after=24)
        add_para("ĐỀ TÀI:", bold=True, size=14)
        p_title = add_para("HỆ THỐNG QUẢN LÝ TÀI CHÍNH CÁ NHÂN TÍCH HỢP AI (EXPENSEAI)", bold=True, size=18, space_after=48)
        p_title.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        
        add_para("Chuyên ngành           : Công nghệ Thông tin", align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=3.0)
        add_para("Giảng viên hướng dẫn : TS. Nguyễn Văn A", align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=3.0)
        add_para("Sinh viên thực hiện   : Lương Dạ Thảo", align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=3.0)
        add_para("Mã sinh viên            : DTC21500000", align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=3.0)
        add_para("Lớp                       : CNTT K20", align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=3.0, space_after=48)
        
        p_footer = add_para("THÁI NGUYÊN, 2026", bold=True, size=13)
        run_break = p_footer.add_run()
        run_break.add_break(WD_BREAK.PAGE)
        print("Cover page added.")
        
    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith('Hình ') or text.startswith('Bảng '):
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                
        if text.startswith('Hình 2.1a'):
            img_path = os.path.join(img_dir, 'Hình 2.1a.png')
            if os.path.exists(img_path):
                new_p = p.insert_paragraph_before()
                new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = new_p.add_run()
                run.add_picture(img_path, width=Inches(6.0))
                print("Inserted 2.1a")
                
        if '[CHÈN SƠ ĐỒ MỚI TẠI ĐÂY]' in text:
            for run in p.runs:
                if '[CHÈN SƠ ĐỒ MỚI TẠI ĐÂY]' in run.text:
                    run.text = run.text.replace('[CHÈN SƠ ĐỒ MỚI TẠI ĐÂY]', '')

    doc.save(out_path)
    print("Done! Saved to", out_path)
except Exception as e:
    print('Error:', e)
