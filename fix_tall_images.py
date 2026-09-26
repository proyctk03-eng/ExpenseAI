import urllib.request
import base64
import zlib
import os
import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

diagrams = {
    'Seq_Cat': '''sequenceDiagram
    autonumber
    actor U as Người dùng
    participant UI as Web/App
    participant API as FastAPI
    participant LLM as Gemini AI
    participant DB as PostgreSQL

    U->>UI: Nhập mô tả (vd: "Ăn phở 50k")
    UI->>API: POST /api/transactions
    API->>API: Xác thực JWT & Rate Limit
    API->>LLM: Gọi AI (Phân loại giao dịch)
    activate LLM
    LLM-->>API: Trả JSON {category, amount}
    deactivate LLM
    API->>DB: Lưu giao dịch vào DB
    activate DB
    DB-->>API: Xác nhận lưu thành công
    deactivate DB
    API-->>UI: Trả về kết quả 200 OK
    UI-->>U: Hiển thị trên Dashboard''',

    'Seq_Adv': '''sequenceDiagram
    autonumber
    actor U as Người dùng
    participant UI as Web/App
    participant API as FastAPI
    participant DB as PostgreSQL
    participant LLM as Gemini AI

    U->>UI: Yêu cầu "Nhận tư vấn"
    UI->>API: GET /api/advice
    API->>DB: Truy vấn dữ liệu thu chi
    activate DB
    DB-->>API: Lịch sử giao dịch 30 ngày
    deactivate DB
    API->>API: Tổng hợp (Income, Expense)
    API->>LLM: Gửi Prompt (Dữ liệu tài chính)
    activate LLM
    LLM-->>API: Trả về tư vấn (Markdown)
    deactivate LLM
    API->>DB: Lưu lịch sử tư vấn
    API-->>UI: Trả về dữ liệu tư vấn
    UI-->>U: Hiển thị Modal lời khuyên''',
    
    'DFD': '''flowchart TD
    classDef ext fill:#F8FAFC,stroke:#334155,stroke-width:2px,color:#0F172A
    classDef proc fill:#F0F7FF,stroke:#2563EB,stroke-width:2px,color:#0F172A,rx:20,ry:20
    classDef ds fill:#ECFDF5,stroke:#059669,stroke-width:2px,color:#0F172A

    U[Người dùng]:::ext
    AI[Google Gemini]:::ext
    
    P1(1.0 Quản lý\nTài khoản):::proc
    P2(2.0 Xử lý\nGiao dịch):::proc
    P3(3.0 Báo cáo\n& Tư vấn):::proc
    
    D1[(D1: Người dùng)]:::ds
    D2[(D2: Giao dịch)]:::ds
    D3[(D3: Danh mục)]:::ds
    
    U -->|Thông tin Đăng nhập| P1
    P1 <-->|Xác thực| D1
    
    U -->|Nhập khoản chi/thu| P2
    P2 -->|Gửi dữ liệu text| AI
    AI -->|Trả về Category, Amount| P2
    P2 -->|Lưu| D2
    P2 -->|Tham chiếu| D3
    
    U -->|Yêu cầu xem| P3
    P3 <-->|Truy vấn| D2
    P3 -->|Gửi Data| AI
    AI -->|Trả về Insight| P3
    P3 -->|Hiển thị| U'''
}

img_dir = r'C:\Users\dathao\Desktop\ảnh mới vẽ'
os.makedirs(img_dir, exist_ok=True)

file_paths = {}
for name, code in diagrams.items():
    compressed = zlib.compress(code.encode('utf-8'))
    b64 = base64.urlsafe_b64encode(compressed).decode('utf-8')
    url = 'https://kroki.io/mermaid/png/' + b64
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    try:
        with urllib.request.urlopen(req, timeout=15) as response:
            file_path = os.path.join(img_dir, f'new_{name}.png')
            with open(file_path, 'wb') as f:
                f.write(response.read())
            file_paths[name] = file_path
        print(f"Generated {name}.png")
    except Exception as e:
        print(f"Failed to generate {name}: {e}")

doc_path = r'C:\Users\dathao\Desktop\nhom02_FINAL_backup_before_videofix.docx'
doc = docx.Document(doc_path)

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip().lower()
    img_to_insert = None
    
    if text.startswith('hình 2.3:'):
        img_to_insert = file_paths.get('Seq_Cat')
        print("Found Hình 2.3")
    elif text.startswith('hình 2.4:'):
        img_to_insert = file_paths.get('Seq_Adv')
        print("Found Hình 2.4")
    elif text.startswith('hình 2.7:'):
        img_to_insert = file_paths.get('DFD')
        print("Found Hình 2.7")
        
    if img_to_insert:
        # Check previous paragraphs and remove old images to clean up the tall ones
        # We look up to 3 paragraphs before the caption to find and clear the one with the image
        for j in range(max(0, i-3), i):
            prev_p = doc.paragraphs[j]
            if '<w:drawing>' in prev_p._element.xml:
                prev_p.clear() # Clear the old image paragraph
                print("Cleared old image.")
                
        # Insert the new proportional image
        new_p = p.insert_paragraph_before()
        new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = new_p.add_run()
        run.add_picture(img_to_insert, width=Inches(6.0))
        print("Inserted new image.")

doc.save(doc_path)
print('Done saving DOCX.')
