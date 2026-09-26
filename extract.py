import docx
import re

doc = docx.Document('docs/Bao_Cao_Hoan_Chinh_ExpenseAI.docx')
matches = set()
pattern = re.compile(r' \([A-Za-z\s-]+\)')

def process_paragraphs(paragraphs):
    for p in paragraphs:
        for match in pattern.findall(p.text):
            matches.add(match)

process_paragraphs(doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            process_paragraphs(cell.paragraphs)

with open('matches.txt', 'w', encoding='utf-8') as f:
    for m in sorted(matches):
        f.write(m + '\n')
