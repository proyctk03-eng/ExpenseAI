import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc_path = r'C:\Users\dathao\Desktop\Bao_Cao_Hoan_Chinh_ExpenseAI.docx'
try:
    doc = docx.Document(doc_path)
    
    # 4. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_h = p_header.add_run("CHƯƠNG 4: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
    run_h.bold = True
    run_h.font.size = Pt(14)
    run_h.font.name = 'Times New Roman'
    p_header.paragraph_format.space_before = Pt(12)
    p_header.paragraph_format.space_after = Pt(6)

    def add_section(title, points):
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_t = p_title.add_run(title)
        run_t.bold = True
        run_t.font.size = Pt(13)
        run_t.font.name = 'Times New Roman'
        p_title.paragraph_format.space_before = Pt(6)
        p_title.paragraph_format.space_after = Pt(6)
        p_title.paragraph_format.line_spacing = 1.15

        for point in points:
            p = doc.add_paragraph(style='List Paragraph')
            run = p.add_run(point)
            run.font.size = Pt(13)
            run.font.name = 'Times New Roman'
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(4)

    # 4.1
    add_section("4.1. Kết quả đạt được", [
        "Về mặt kỹ thuật: Triển khai thành công kiến trúc Client-Server hiện đại với FastAPI (Backend) và cơ sở dữ liệu PostgreSQL. Hệ thống có khả năng xử lý bất đồng bộ, chịu tải tốt nhờ tích hợp Redis Cache.",
        "Về mặt ứng dụng AI: Tích hợp thành công các luồng AI thực tiễn bao gồm: tự động nhận diện danh mục chi tiêu, trích xuất dữ liệu hóa đơn (OCR), và trợ lý AI tư vấn tài chính cá nhân hóa.",
        "Về mặt bảo mật: Hoàn thiện các tiêu chuẩn bảo mật cơ bản như xác thực JWT, phân quyền Role-based, mã hóa mật khẩu và cấu hình an toàn cho HTTP Cookies/Headers."
    ])

    # 4.2
    add_section("4.2. Hạn chế của hệ thống", [
        "Phụ thuộc vào API của bên thứ ba (OpenAI/Gemini). Khi có sự cố kết nối mạng hoặc giới hạn hạn mức, tính năng AI sẽ bị gián đoạn.",
        "Tốc độ xử lý ảnh hóa đơn (OCR) đôi khi còn chậm do phải chờ phản hồi từ model ngôn ngữ lớn.",
        "Hệ thống hiện tại chỉ hoạt động trên nền tảng Web, chưa có phiên bản Native Mobile App để người dùng trải nghiệm tiện lợi hơn trên điện thoại."
    ])

    # 4.3
    add_section("4.3. Hướng phát triển", [
        "Nâng cấp đa nền tảng: Phát triển thêm ứng dụng di động (Mobile App) bằng React Native hoặc Flutter để tăng trải nghiệm người dùng.",
        "Tối ưu hóa AI: Nghiên cứu nhúng các mô hình AI nhỏ gọn (Local LLM) chạy trực tiếp trên server để giảm sự phụ thuộc và chi phí cho API bên thứ ba.",
        "Tích hợp Open Banking: Liên kết trực tiếp với API của các ngân hàng và ví điện tử để đồng bộ hóa lịch sử giao dịch tự động."
    ])

    doc.save(doc_path)
    print("Thêm thành công nội dung vào file Word.")
except Exception as e:
    print("Lỗi:", e)
