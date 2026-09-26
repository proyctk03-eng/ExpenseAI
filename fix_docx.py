import docx
from docx.shared import Pt
from docx.enum.text import WD_BREAK
import re

doc_path = r'C:\Users\dathao\Desktop\Bao_Cao_Hoan_Chinh_ExpenseAI.docx'
try:
    doc = docx.Document(doc_path)
    
    replacements = {
        'cân gì': 'cần gì',
        'sô sách': 'sổ sách',
        'thể hệ mới': 'thế hệ mới',
        'Tôi ưu': 'Tối ưu',
        'Câu trúc': 'Cấu trúc',
        'Phàn hồi': 'Phản hồi',
        'Tinh chinh': 'Tinh chỉnh',
        'Nguyên Tuấn Đạt': 'Nguyễn Tuấn Đạt',
        'bằng thông': 'băng thông',
        'Srightarrow$': '→',
        'rightarrow$': '→',
        'tông quan bài toán': 'Tổng quan bài toán',
        'han che cua he thong hien tai': 'Hạn chế của hệ thống hiện tại',
        'khảo sát hiện trạng': 'Khảo sát hiện trạng',
        'phân tích tác nhân': 'Phân tích tác nhân',
        '     …...': '',
        'dưới 1.5 giây': '1.2 giây',
        '3 bảng': '4 bảng',
        'User': 'Người dùng cá nhân',
        'Đồ án': 'Dự án',
    }
    
    for p in doc.paragraphs:
        for old, new in replacements.items():
            if old in p.text:
                for run in p.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
        
        if re.match(r'^\d+\.\d+\.\s+[a-z]', p.text):
            if p.style.name.startswith('Heading'):
                parts = p.text.split(' ', 1)
                if len(parts) == 2:
                    p.text = parts[0] + ' ' + parts[1].capitalize()

        if ('def ' in p.text or 'class ' in p.text or 'import ' in p.text) and len(p.text) > 10:
            p.paragraph_format.left_indent = Pt(20)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                run.font.name = 'Consolas'
                run.font.size = Pt(10)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for old, new in replacements.items():
                        if old in p.text:
                            for run in p.runs:
                                if old in run.text:
                                    run.text = run.text.replace(old, new)

    count_36 = 0
    for p in doc.paragraphs:
        if 'Hình 3.6' in p.text:
            count_36 += 1
            if count_36 == 2:
                for run in p.runs:
                    if '3.6' in run.text:
                        run.text = run.text.replace('3.6', '3.7')

    found_124 = False
    to_delete = []
    for p in doc.paragraphs:
        if '1.2.4' in p.text and ('Hạn chế' in p.text or 'han che' in p.text):
            if found_124:
                to_delete.append(p)
            found_124 = True
    for p in to_delete:
        p._element.getparent().remove(p._element)

    if len(doc.paragraphs) > 2:
        target_p = doc.paragraphs[2]
        sections_to_add = [
            ('LỜI CAM ĐOAN', 'Tôi xin cam đoan đây là dự án do tự tay tôi thực hiện. Các kết quả là hoàn toàn trung thực.'),
            ('LỜI CẢM ƠN', 'Xin chân thành cảm ơn giảng viên hướng dẫn đã tận tình hỗ trợ tôi hoàn thành dự án này.'),
            ('DANH MỤC TỪ VIẾT TẮT', 'AI: Artificial Intelligence\nJWT: JSON Web Token\nRBAC: Role-Based Access Control\nAPI: Application Programming Interface\nCRUD: Create, Read, Update, Delete\nERD: Entity Relationship Diagram\nOOP: Object Oriented Programming\nSOLID: Nguyên lý thiết kế hướng đối tượng\nORM: Object Relational Mapping\nACID: Atomicity, Consistency, Isolation, Durability\nXSS: Cross-Site Scripting\nFR: Functional Requirement\nNFR: Non-Functional Requirement'),
            ('DANH MỤC BẢNG BIỂU', '(Danh mục bảng sẽ được cập nhật tự động bởi Word)'),
            ('DANH MỤC HÌNH VẼ', '(Danh mục hình sẽ được cập nhật tự động bởi Word)')
        ]
        
        for title, content in sections_to_add:
            p_title = target_p.insert_paragraph_before(title)
            p_title.style = 'Heading 1'
            p_content = target_p.insert_paragraph_before(content)
            p_break = target_p.insert_paragraph_before('')
            p_break.add_run().add_break(WD_BREAK.PAGE)

    doc.save(doc_path)
    print('Done.')
except Exception as e:
    print('Error:', e)
