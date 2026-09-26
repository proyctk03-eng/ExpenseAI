import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc_path = r'C:\Users\dathao\Desktop\Bao_Cao_Hoan_Chinh_ExpenseAI.docx'
try:
    doc = docx.Document(doc_path)
    
    count = 0
    for p in doc.paragraphs:
        if '<w:drawing' in p._p.xml:
            p.clear()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            
            run = p.add_run("[CHÈN SƠ ĐỒ MỚI TẠI ĐÂY]")
            run.bold = True
            run.font.color.rgb = RGBColor(255, 0, 0)
            count += 1
            
    # Also check inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if '<w:drawing' in p._p.xml:
                        p.clear()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(0)
                        p.paragraph_format.line_spacing = 1.0
                        run = p.add_run("[CHÈN SƠ ĐỒ MỚI TẠI ĐÂY]")
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 0, 0)
                        count += 1

    doc.save(doc_path)
    print(f"Done. Replaced {count} images with tight placeholders.")
except Exception as e:
    print('Error:', e)
