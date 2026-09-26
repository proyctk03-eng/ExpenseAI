import docx
import re
from docx.shared import RGBColor
from docx.oxml.ns import qn

def fix_heading(text):
    if not text.strip(): return text
    if text.isupper(): return text # Giữ nguyên nếu viết hoa toàn bộ (CHƯƠNG 1...)
    
    # Tách từ để xử lý Title Case -> Sentence case
    # VD: "Tổng Quan Về Ứng Dụng" -> "Tổng quan về ứng dụng"
    tokens = re.split(r'(\s+)', text)
    new_tokens = []
    
    first_word_done = False
    for t in tokens:
        if t.strip() == '':
            new_tokens.append(t)
            continue
            
        if t in ["AI", "ExpenseAI", "API", "UI", "UX", "H&M", "ICTU", "ERD", "OOP", "RBAC", "VNĐ", "VND", "USD"]:
            new_tokens.append(t)
            first_word_done = True
        elif not first_word_done:
            new_tokens.append(t)
            first_word_done = True
        else:
            new_tokens.append(t.lower())
            
    return ''.join(new_tokens)

def fix_quotes(text):
    # Thay ngoặc kép thẳng thành ngoặc kép cong
    text = re.sub(r'"([^"]*)"', r'“\1”', text)
    return text

def process_doc():
    doc = docx.Document('docs/Bao_Cao_Hoan_Chinh_ExpenseAI.docx')
    
    # Xử lý đoạn văn
    for p in doc.paragraphs:
        is_heading = p.style.name.startswith('Heading')
        
        # Sửa Heading casing
        if is_heading:
            old_text = p.text
            new_text = fix_heading(old_text)
            if new_text != old_text:
                # Đối với Heading, an toàn nhất là xóa hết run và gán lại text, 
                # vì Heading style tự lo phần in đậm/font chữ
                p.text = new_text
        
        for run in p.runs:
            if not run.text: continue
            
            # Sửa dấu gạch nối dài thành gạch ngắn
            run.text = run.text.replace('—', '-').replace('–', '-')
            
            # Sửa dấu bullet point thành gạch ngang
            run.text = run.text.replace('•', '-')
            
            # Sửa ngoặc kép
            run.text = fix_quotes(run.text)

    # Xử lý Bảng biểu (Đưa về đen trắng)
    for table in doc.tables:
        table.style = 'Table Grid' # Chuyển về bảng đen trắng cơ bản
        for row in table.rows:
            for cell in row.cells:
                # Xóa màu nền của ô (Shading)
                tcPr = cell._tc.get_or_add_tcPr()
                shd = tcPr.find(qn('w:shd'))
                if shd is not None:
                    tcPr.remove(shd)
                    
                # Xử lý text trong bảng
                for p in cell.paragraphs:
                    for run in p.runs:
                        if run.text:
                            run.font.color.rgb = RGBColor(0, 0, 0) # Ép màu đen
                            run.text = run.text.replace('—', '-').replace('–', '-')
                            run.text = run.text.replace('•', '-')
                            run.text = fix_quotes(run.text)

    # Ghi đè file
    doc.save('docs/Bao_Cao_Hoan_Chinh_ExpenseAI.docx')
    print("Done de-AIing the document.")

if __name__ == '__main__':
    process_doc()
