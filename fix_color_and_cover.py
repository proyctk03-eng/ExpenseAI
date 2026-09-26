import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

doc_path = r'C:\Users\dathao\Desktop\Bao_Cao_Hoan_Chinh_ExpenseAI.docx'
try:
    doc = docx.Document(doc_path)
    
    # 1. Fix blue text for captions
    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith('Hình ') or text.startswith('Bảng '):
            for run in p.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    text = p.text.strip()
                    if text.startswith('Hình ') or text.startswith('Bảng '):
                        for run in p.runs:
                            run.font.color.rgb = RGBColor(0, 0, 0)
                            
    # 2. Add Cover Page
    # Check if cover page already added to avoid duplicates
    if doc.paragraphs[0].text != 'ĐẠI HỌC THÁI NGUYÊN':
        for section in doc.sections:
            section.top_margin = Inches(2.0 / 2.54)
            section.bottom_margin = Inches(2.0 / 2.54)
            section.left_margin = Inches(3.5 / 2.54)
            section.right_margin = Inches(2.0 / 2.54)
            
        first_p = doc.paragraphs[0]
        
        def add_para(text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=False, size=13, space_after=0, left_indent=0):
            new_p = first_p.insert_paragraph_before('')
            new_p.alignment = align
            new_p.paragraph_format.space_after = Pt(space_after)
            if left_indent > 0:
                new_p.paragraph_format.left_indent = Inches(left_indent / 2.54)
            run = new_p.add_run(text)
            run.bold = bold
            run.font.size = Pt(size)
            run.font.name = 'Times New Roman'
            run.font.color.rgb = RGBColor(0, 0, 0)
            return new_p

        add_para("ĐẠI HỌC THÁI NGUYÊN", bold=False, size=13)
        add_para("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG", bold=True, size=13)
        add_para("------------------***------------------", bold=False, size=11, space_after=36)
        
        add_para("BÁO CÁO HỌC PHẦN / BÀI TẬP LỚN", bold=True, size=16, space_after=24)
        
        add_para("ĐỀ TÀI:", bold=True, size=14)
        p_title = add_para("HỆ THỐNG QUẢN LÝ TÀI CHÍNH CÁ NHÂN TÍCH HỢP AI (EXPENSEAI)", bold=True, size=18, space_after=48)
        p_title.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        
        add_para("Chuyên ngành           : [Điền chuyên ngành của bạn]", align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=3.0)
        add_para("Giảng viên hướng dẫn : [Điền tên Giảng viên]", align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=3.0)
        add_para("Sinh viên thực hiện   : Lương Dạ Thảo", align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=3.0)
        add_para("Mã sinh viên            : [Điền Mã SV]", align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=3.0)
        add_para("Lớp                       : [Điền Lớp]", align=WD_ALIGN_PARAGRAPH.LEFT, left_indent=3.0, space_after=48)
        
        p_footer = add_para("THÁI NGUYÊN, 2026", bold=True, size=13)
        
        run_break = p_footer.add_run()
        run_break.add_break(WD_BREAK.PAGE)

    doc.save(doc_path)
    print("Fixes and cover page applied successfully.")
except Exception as e:
    print('Error:', e)
