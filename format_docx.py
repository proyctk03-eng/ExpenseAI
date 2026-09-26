import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc_path = r'C:\Users\dathao\Desktop\Bao_Cao_Hoan_Chinh_ExpenseAI.docx'
try:
    doc = docx.Document(doc_path)
    
    replacements = {
        "(chia nhỏ sơ đồ)": "",
        "Người dùng cá nhân (Người dùng)": "Người dùng cá nhân",
        "Role (Vai trò)": "Role",
        "Permission (Quyền hạn)": "Permission",
        "Category (Danh mục)": "Category",
        "Transaction (Giao dịch)": "Transaction",
        "Budget (Ngân sách)": "Budget",
        "AIPrediction (Dự đoán AI)": "AIPrediction",
        "(Khóa chính)": "(PK)"
    }
    
    def process_paragraph(p):
        for old, new in replacements.items():
            if old in p.text:
                for run in p.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
        
        is_heading = p.style.name.startswith('Heading') or p.style.name.startswith('Title')
        
        is_code = False
        if p.runs and p.runs[0].font.name == 'Consolas':
            is_code = True
            
        if not is_heading and not is_code and p.text.strip():
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            # p.paragraph_format.line_spacing = 1.3 # Avoid overriding tight spacing for placeholders
            for run in p.runs:
                if run.font.name != 'Consolas':
                    run.font.name = 'Times New Roman'
                    if run.font.size != Pt(10): # Don't override code or caption sizes if possible, but ICTU is strict.
                        run.font.size = Pt(13)

    for p in doc.paragraphs:
        process_paragraph(p)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_paragraph(p)
                    
    doc.save(doc_path)
    print("Format applied successfully.")
except Exception as e:
    print("Error:", e)
