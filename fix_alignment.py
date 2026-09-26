import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc_path = r'C:\Users\dathao\Desktop\Bao_Cao_Hoan_Chinh_ExpenseAI.docx'
try:
    doc = docx.Document(doc_path)
    
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
            
        is_code = False
        if p.runs and p.runs[0].font.name == 'Consolas':
            is_code = True
        
        is_caption = text.startswith('Hình ') or text.startswith('Bảng ') or 'CHÈN SƠ ĐỒ' in text
        
        is_heading = p.style.name.startswith('Heading')
        is_chapter = 'CHƯƠNG ' in text.upper()
        
        if is_code:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif is_caption:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif is_chapter:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif is_heading:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.runs and p.runs[0].font.name == 'Consolas':
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        
    doc.save(doc_path)
    print("Alignment fixed.")
except Exception as e:
    print('Error:', e)
