import docx
doc_path = r'C:\Users\dathao\Desktop\Bao_Cao_Final_ExpenseAI.docx'
try:
    doc = docx.Document(doc_path)
    with open('check_cover.txt', 'w', encoding='utf-8') as f:
        for i in range(min(15, len(doc.paragraphs))):
            f.write(f"[{i}] {doc.paragraphs[i].text.strip()}\n")
except Exception as e:
    print('Error:', e)
