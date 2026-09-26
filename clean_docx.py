import docx
from docx.shared import Pt, Inches
import re

doc_path = r'C:\Users\dathao\Desktop\Bao_Cao_Hoan_Chinh_ExpenseAI.docx'
try:
    doc = docx.Document(doc_path)
    
    empty_paragraphs = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            empty_paragraphs.append(p)
            continue
        
        if text.startswith('Tóm lại,') or text.startswith('Nhìn chung,') or text.startswith('Có thể thấy,'):
            empty_paragraphs.append(p)
            continue
            
        for run in p.runs:
            if '–' in run.text or '—' in run.text:
                run.text = run.text.replace('–', '-').replace('—', '-')
                
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.space_before = Pt(0)

    def delete_paragraph(paragraph):
        p = paragraph._element
        if p.getparent() is not None:
            p.getparent().remove(p)
            paragraph._p = paragraph._element = None

    for p in empty_paragraphs:
        try:
            delete_paragraph(p)
        except:
            pass

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                shading_elms = tcPr.xpath('w:shd')
                if shading_elms:
                    for shd in shading_elms:
                        tcPr.remove(shd)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.color.rgb = docx.shared.RGBColor(0, 0, 0)

    for shape in doc.inline_shapes:
        if shape.width:
            shape.width = int(shape.width * 0.8)
        if shape.height:
            shape.height = int(shape.height * 0.8)

    doc.save(doc_path)
    print('Docx cleaned and compressed successfully.')
except Exception as e:
    print('Error:', e)
