import docx
import re
doc = docx.Document('docs/Bao_Cao_Hoan_Chinh_ExpenseAI.docx')
pattern = re.compile(r' \([A-Za-z\s-]+\)')
split_runs_count = 0

def check(paragraphs):
    global split_runs_count
    for p in paragraphs:
        if pattern.search(p.text):
            found_in_run = False
            for run in p.runs:
                if pattern.search(run.text) or re.search(r'\([A-Za-z\s-]+\)', run.text):
                    found_in_run = True
            if not found_in_run:
                split_runs_count += 1

check(doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            check(cell.paragraphs)

print("Split runs count:", split_runs_count)
