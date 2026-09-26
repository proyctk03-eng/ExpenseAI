import urllib.request
import base64
import zlib
import json
import os
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 1. Define diagrams
diagrams = {
    '1.1a': '''flowchart TD
    A([Bắt đầu]) --> B[Mua sắm/Thanh toán]
    B --> C{Có nhớ ghi lại ngay?}
    C -- Không --> D[Quên khoản chi, sai số liệu]
    C -- Có --> E[Mở sổ tay / Excel]
    E --> F[Tính toán cộng trừ bằng tay]
    F --> G[Ghi chép thủ công]
    D --> H([Kết thúc: Dữ liệu sai lệch])
    G --> I([Kết thúc: Mất thời gian, chán nản])''',
    
    '1.1b': '''flowchart TD
    A([Bắt đầu]) --> B[Người dùng nhập thoại/text tự do]
    B --> C[AI bóc tách số tiền & danh mục]
    C --> D{AI Phân loại thành công?}
    D -- Không --> E[Yêu cầu nhập lại/Sửa lỗi]
    D -- Có --> F[Ghi nhận vào CSDL tự động]
    F --> G[Kiểm tra ngân sách (Budget)]
    G --> H([Kết thúc: Dữ liệu chuẩn xác, nhanh chóng])
    E --> B''',
    
    '2.1a': '''flowchart LR
    User((Người Dùng))
    Gemini((Google Gemini AI))
    User --> UC1([Đăng ký/Đăng nhập])
    User --> UC2([Quản lý Giao dịch])
    User --> UC3([Thiết lập Ngân sách])
    User --> UC4([Nhận Tư vấn Tài chính])
    UC2 -.->|Gọi API| Gemini
    UC4 -.->|Phân tích| Gemini''',
    
    '2.2a': '''flowchart TD
    Start([Nhận Request]) --> Guard[Kiểm tra bảo mật & JWT]
    Guard --> CheckAI{Yêu cầu dùng AI?}
    CheckAI -- Không --> Validate[Kiểm thực Pydantic thủ công]
    CheckAI -- Có --> CallAI[Gửi Text tới Gemini API]
    CallAI --> Parse[Bóc tách JSON]
    Parse --> Validate
    Validate --> Save[(Ghi vào PostgreSQL)]
    Save --> CalcBudget[Tính toán hạn mức ngân sách]
    CalcBudget --> Alert{Vượt 80% ngân sách?}
    Alert -- Có --> SendAlert[Gửi cảnh báo Vàng/Đỏ]
    Alert -- Không --> Finish([Trả về HTTP 200 OK])
    SendAlert --> Finish''',
    
    '2.3': '''sequenceDiagram
    actor U as Người dùng
    participant F as Frontend
    participant API as FastAPI Router
    participant S as AI Service
    participant DB as PostgreSQL
    U->>F: Nhập "Lương tháng này 10 củ"
    F->>API: POST /transactions/ai
    API->>S: Chuyển Prompt
    S->>S: Guardrails Check
    S-->>API: Trả JSON
    API->>DB: INSERT Transaction
    DB-->>API: Success
    API-->>F: Trả kết quả hiển thị''',
    
    '2.6': '''erDiagram
    USERS ||--o{ TRANSACTIONS : "thực hiện"
    USERS ||--o{ BUDGETS : "thiết lập"
    CATEGORIES ||--o{ TRANSACTIONS : "phân loại"
    USERS { UUID id PK }
    CATEGORIES { INT id PK }
    TRANSACTIONS { UUID id PK }
    BUDGETS { UUID id PK }''',
    
    '3.1': '''graph TD
    subgraph Client_Tier [Client Tier / Presentation]
        Browser[Trình duyệt Web]
        Mobile[Thiết bị Di động]
    end
    subgraph App_Tier [Application Tier / Business Logic]
        FastAPI[FastAPI Server - Python]
        AI_Module[AI Guardrails Module]
    end
    subgraph Data_Tier [Data Tier / Access]
        Postgres[(PostgreSQL Database)]
        Redis[(Redis Cache)]
    end
    Client_Tier == HTTPS ==> App_Tier
    App_Tier == SQL/TCP ==> Data_Tier''',
    
    '3.2': '''graph TD
    subgraph Deployment
        Client((Client)) --> Nginx[Nginx Reverse Proxy]
        Nginx --> API[FastAPI Docker Container]
        API --> DB[(PostgreSQL Container)]
        API --> External((Google Gemini API))
    end''',
    
    '3.3': '''graph TD
    UI[Frontend Component] --> Router[API Router Component]
    Router --> Auth[Auth Service]
    Router --> Transaction[Transaction Service]
    Transaction --> AI[Gemini API Client]
    Transaction --> ORM[SQLAlchemy ORM]
    ORM --> DB[(PostgreSQL)]''',
    
    '3.4': '''classDiagram
    class BaseService {
        +db: Session
        +get(id)
        +create(obj)
    }
    class TransactionService {
        +analyze_with_ai(text)
    }
    class AuthService {
        +verify_password(plain, hashed)
    }
    BaseService <|-- TransactionService
    BaseService <|-- AuthService''',
    
    '3.5a': '''stateDiagram-v2
    [*] --> Nhap_Lieu
    Nhap_Lieu --> Tien_Xu_Ly : Bấm Gửi
    Tien_Xu_Ly --> Goi_API_AI : Dữ liệu hợp lệ
    Tien_Xu_Ly --> Bao_Loi : Dữ liệu sai
    Goi_API_AI --> Cho_Tra_Loi
    Cho_Tra_Loi --> Ghi_Nhan_Thanh_Cong : JSON chuẩn xác
    Cho_Tra_Loi --> Fallback : AI Ảo giác / Mất mạng
    Fallback --> Ghi_Nhan_Thanh_Cong : Chuyển qua thủ công
    Ghi_Nhan_Thanh_Cong --> [*]'''
}

img_dir = 'C:/Users/dathao/Downloads/AI/ExpenseAI/scratch/images'
os.makedirs(img_dir, exist_ok=True)

for name, code in diagrams.items():
    try:
        compressed = zlib.compress(code.encode('utf-8'))
        b64 = base64.urlsafe_b64encode(compressed).decode('utf-8')
        url = 'https://kroki.io/mermaid/png/' + b64
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response:
            with open(os.path.join(img_dir, f'Hình {name}.png'), 'wb') as f:
                f.write(response.read())
        print(f"Generated {name}")
    except Exception as e:
        print(f"Failed to generate {name}: {e}")

# 2. Insert into DOCX
doc_path = r'C:\Users\dathao\Desktop\Bao_Cao_Hoan_Chinh_ExpenseAI.docx'
out_path = r'C:\Users\dathao\Desktop\Bao_Cao_Final_ExpenseAI.docx'
try:
    doc = docx.Document(doc_path)
    
    for p in doc.paragraphs:
        text = p.text.strip()
        if text.startswith('Hình '):
            # Extract number
            parts = text.split(':')
            if len(parts) > 0:
                fig_name = parts[0].strip() # e.g. "Hình 1.1a"
                img_path = os.path.join(img_dir, f'{fig_name}.png')
                if os.path.exists(img_path):
                    # Insert image before caption
                    new_p = p.insert_paragraph_before()
                    new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = new_p.add_run()
                    run.add_picture(img_path, width=Inches(6.0))
                    print(f"Inserted {fig_name}")
                    
        # Also clean up placeholder if any
        if '[CHÈN SƠ ĐỒ MỚI TẠI ĐÂY]' in p.text:
            p.text = p.text.replace('[CHÈN SƠ ĐỒ MỚI TẠI ĐÂY]', '')

    doc.save(out_path)
    print("All images inserted. Saved as Bao_Cao_Final_ExpenseAI.docx")
except Exception as e:
    print('Error:', e)
